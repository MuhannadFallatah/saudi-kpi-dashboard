# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import threading
import time
import calendar
import re
import ssl
from copy import copy
from io import BytesIO
from datetime import date, datetime, timedelta, timezone

RIYADH_TZ = timezone(timedelta(hours=3), "Asia/Riyadh")
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


APP_DIR = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = APP_DIR / "Saudi_16_KPI_Data_Mart.xlsx"
WORKBOOK_PATH = Path(os.environ.get("KPI_DATA_WORKBOOK", DEFAULT_WORKBOOK))
REFRESH_MINUTES = max(1, int(os.environ.get("KPI_REFRESH_MINUTES", "1440")))
HOST = "127.0.0.1"
PORT = int(os.environ.get("KPI_PORT", "8765"))
ALLOWED_HOSTS = {"www.stats.gov.sa", "www.mof.gov.sa", "api.datasaudi.sa"}
DATASAUDI_PMI_URL = (
    "https://api.datasaudi.sa/tesseract/data.jsonrecords"
    "?cube=pmi&drilldowns=Month&measures=Purchasing+Manager+Index"
)
DATASAUDI_PMI_PAGE = "https://datasaudi.sa/en/#purchasing-manager-index"

FACT_HEADERS = (
    "CountryCode",
    "KPI_ID",
    "PeriodStart",
    "PeriodEnd",
    "Year",
    "Month",
    "MonthName",
    "Value",
    "Unit",
    "SourceTable",
    "ObservationStatus",
    "SourceURL",
    "SourceNote",
    "Open",
    "High",
    "Low",
    "Close",
    "VolumeTraded",
    "TurnoverSAR",
    "TradeCount",
)
BEC_CODES = {"112", "212", "312", "412", "512", "612", "712", "812"}
KPI_NAMES_EN = {
    "KPI_01": "Purchasing Managers' Index (PMI)",
    "KPI_02": "Point-of-sale transactions (POS)",
    "KPI_03": "Cement sales",
    "KPI_04": "Private-sector bank credit (12a)",
    "KPI_05": "Total goods imports",
    "KPI_06": "Private-sector imports",
    "KPI_07": "Capital imports",
    "KPI_08": "Non-oil exports including re-exports",
    "KPI_09": "Monthly average SAIBOR",
    "KPI_10": "Investment licences",
    "KPI_11": "Oil exports",
    "KPI_12": "Goods exports",
    "KPI_13": "Issued building permits",
    "KPI_14": "Credit default swap (CDS) spreads",
    "KPI_15": "Tadawul All Share Index (TASI)",
    "KPI_16": "Government capital expenditure",
}
KPI_SHEET_SLUGS = {
    "KPI_01": "PMI",
    "KPI_02": "POS_Transactions",
    "KPI_03": "Cement_Sales",
    "KPI_04": "Private_Bank_Credit",
    "KPI_05": "Total_Goods_Imports",
    "KPI_06": "Private_Sector_Imports",
    "KPI_07": "Capital_Imports",
    "KPI_08": "NonOil_Exports",
    "KPI_09": "SAIBOR",
    "KPI_10": "Investment_Licences",
    "KPI_11": "Oil_Exports",
    "KPI_12": "Goods_Exports",
    "KPI_13": "Building_Permits",
    "KPI_14": "CDS_Spreads",
    "KPI_15": "TASI",
    "KPI_16": "Govt_Capex",
}
MONTHS = {
    name.lower(): number
    for number, name in enumerate(
        ("", "January", "February", "March", "April", "May", "June", "July",
         "August", "September", "October", "November", "December")
    )
    if name
}


def serialize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def workbook_data(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Data workbook does not exist: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"KPI_Catalog", "Source_Register", "Fact_All"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        workbook.close()
        raise ValueError(f"Workbook is missing required sheets: {', '.join(sorted(missing))}")

    catalog_sheet = workbook["KPI_Catalog"]
    catalog_headers = [cell.value for cell in catalog_sheet[1]]
    catalog = [
        dict(zip(catalog_headers, row))
        for row in catalog_sheet.iter_rows(min_row=2, values_only=True)
        if row and row[0]
    ]
    source_sheet = workbook["Source_Register"]
    source_headers = [cell.value for cell in source_sheet[1]]
    sources = {
        row[0]: dict(zip(source_headers, row))
        for row in source_sheet.iter_rows(min_row=2, values_only=True)
        if row and row[0]
    }

    facts = []
    for row in workbook["Fact_All"].iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        facts.append(dict(zip(FACT_HEADERS, (serialize(value) for value in row[: len(FACT_HEADERS)]))))
    components = []
    if "Fact_KPI07_BEC" in workbook.sheetnames:
        component_sheet = workbook["Fact_KPI07_BEC"]
        component_headers = [cell.value for cell in component_sheet[1]]
        components = [
            dict(zip(component_headers, (serialize(value) for value in row)))
            for row in component_sheet.iter_rows(min_row=2, values_only=True)
            if row and row[1] == "KPI_07"
        ]
    export_components = []
    if "08_KPI_Components" in workbook.sheetnames:
        component_sheet = workbook["08_KPI_Components"]
        component_headers = [cell.value for cell in component_sheet[1]]
        export_components = [
            dict(zip(component_headers, (serialize(value) for value in row)))
            for row in component_sheet.iter_rows(min_row=2, values_only=True)
            if row and row[1] == "KPI_08"
        ]
    workbook.close()
    return {
        "catalog": catalog,
        "sources": sources,
        "facts": facts,
        "components": components,
        "kpi08_components": export_components,
    }


def tls_context() -> ssl.SSLContext:
    context = ssl.create_default_context()
    # www.mof.gov.sa omits its DigiCert intermediate certificate from the TLS handshake.
    for pem in sorted((APP_DIR / "certs").glob("*.pem")):
        context.load_verify_locations(cafile=str(pem))
    return context


def download_bytes(url: str) -> bytes:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in ALLOWED_HOSTS:
        raise ValueError("Live extractors only allow HTTPS URLs on approved source hosts.")
    request = Request(url, headers={"User-Agent": "Saudi-KPI-Monitor/1.0"})
    with urlopen(request, timeout=45, context=tls_context()) as response:
        if response.status != 200:
            raise OSError(f"Source returned HTTP {response.status}: {url}")
        return response.read()


def make_quarter_fact(
    kpi_id: str, year: int, quarter: int, value: float, unit: str,
    source_table: str, source_url: str, source_note: str,
) -> dict[str, Any]:
    first_month = 3 * quarter - 2
    fact = make_fact(kpi_id, year, first_month, value, unit, source_table, source_url, source_note)
    last_month = first_month + 2
    fact["PeriodEnd"] = date(year, last_month, calendar.monthrange(year, last_month)[1]).isoformat()
    fact["Month"] = None
    fact["MonthName"] = f"Q{quarter}"
    return fact


MOF_CAPEX_LABEL = "non-financial assets (capex)"
MOF_PERIOD_RE = re.compile(
    r"Actual\s+expenditures?\s+(?:in|until)\s+(Q[1-4]|H1)\s+of\s+FY\s+(\d{4})", re.IGNORECASE
)
MOF_PERIOD_QUARTERS = {"Q1": 1, "H1": 2, "Q2": 2, "Q3": 3, "Q4": 4}


def extract_mof_capex(content: bytes, source_url: str) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        # The current report sheet carries a trailing space; unnamed-suffix copies are stale templates.
        names = sorted(
            (name for name in workbook.sheetnames if name.strip() == "Expenditures"),
            key=lambda name: name == name.strip(),
        )
        if not names:
            raise ValueError("MOF workbook has no Expenditures sheet.")
        rows = list(workbook[names[0]].iter_rows(values_only=True))
    finally:
        workbook.close()
    period = None
    for row in rows[:8]:
        for cell in row:
            match = MOF_PERIOD_RE.search(cell) if isinstance(cell, str) else None
            if match:
                period = match
                break
        if period:
            break
    if not period:
        raise ValueError("MOF workbook period title was not found.")
    quarters = MOF_PERIOD_QUARTERS[period.group(1).upper()]
    year = int(period.group(2))
    values: list[float] | None = None
    for row in rows:
        if any(isinstance(cell, str) and cell.strip().lower() == MOF_CAPEX_LABEL for cell in row):
            values = [float(c) for c in row if isinstance(c, (int, float)) and not isinstance(c, bool)]
            break
    if values is None:
        raise ValueError("MOF CAPEX row was not found.")
    if quarters == 1:
        by_quarter = values[:1]
    else:
        by_quarter = values[1 : quarters + 1]
        if len(by_quarter) != quarters or abs(sum(by_quarter) - values[0]) > 1:
            raise ValueError("MOF CAPEX quarterly values do not reconcile with the cumulative total.")
        by_quarter.reverse()
    if len(by_quarter) != quarters:
        raise ValueError("MOF CAPEX quarterly values are incomplete.")
    return [
        make_quarter_fact(
            "KPI_16", year, quarter, value, "SAR million",
            "MOF Quarterly Budget Performance Report - Expenditures: Non-financial Assets (CAPEX)",
            source_url, "Actual quarterly government capital expenditure; quarter derived from the report's period columns.",
        )
        for quarter, value in enumerate(by_quarter, start=1)
    ]


def mof_report_urls(fetch) -> list[str]:
    # Latest report workbook per fiscal year, discovered from the MOF year pages.
    urls: list[str] = []
    for year in range(2022, date.today().year + 1):
        links: list[str] = []
        last_problem = "no workbook links on the page"
        for slug in (f"budget-performance-report-{year}-1", f"budget-performance-report-{year}"):
            try:
                html = fetch(f"https://www.mof.gov.sa/ar/w/{slug}").decode("utf-8", "replace")
            except OSError as error:
                last_problem = f"{slug}: {error}"
                continue
            links = list(dict.fromkeys(re.findall(r"/documents/[^\"'\s<>]*\.xlsx[^\"'\s<>]*", html)))
            if links:
                break
            last_problem = f"{slug}: no workbook links in {len(html)} characters of HTML"
        if not links:
            if year == date.today().year:
                continue
            raise ValueError(f"No MOF budget performance workbook found for {year} ({last_problem}).")

        def quarter_of(link: str) -> int:
            name = unquote(link.split("/")[4]) if link.count("/") >= 4 else link
            match = re.search(r"Q([1-4])", name)
            return int(match.group(1)) if match else 0

        best = max(links, key=quarter_of)
        urls.append("https://www.mof.gov.sa" + best.replace("?download=true", "").replace("&amp;", "&"))
    return urls


def make_fact(
    kpi_id: str,
    year: int,
    month: int,
    value: float,
    unit: str,
    source_table: str,
    source_url: str,
    source_note: str,
    status: str = "Reported",
) -> dict[str, Any]:
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    return {
        "CountryCode": "SAU",
        "KPI_ID": kpi_id,
        "PeriodStart": start.isoformat(),
        "PeriodEnd": end.isoformat(),
        "Year": year,
        "Month": month,
        "MonthName": start.strftime("%B"),
        "Value": float(value),
        "Unit": unit,
        "SourceTable": source_table,
        "ObservationStatus": status,
        "SourceURL": source_url,
        "SourceNote": source_note,
        "Open": None,
        "High": None,
        "Low": None,
        "Close": None,
        "VolumeTraded": None,
        "TurnoverSAR": None,
        "TradeCount": None,
    }


def extract_datasaudi_pmi(content: bytes) -> list[dict[str, Any]]:
    try:
        payload = json.loads(content)
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError("DataSaudi PMI response is not valid JSON.") from error

    if not isinstance(payload, dict):
        raise ValueError("DataSaudi PMI response must be a JSON object.")

    expected_measure = "Purchasing Manager Index"
    columns = payload.get("columns")
    data = payload.get("data")
    page = payload.get("page")
    total = page.get("total") if isinstance(page, dict) else None
    if not isinstance(columns, list) or expected_measure not in columns:
        raise ValueError("DataSaudi PMI response is missing the expected measure column.")
    if not isinstance(data, list) or not data:
        raise ValueError("DataSaudi PMI response contains no monthly observations.")
    if not isinstance(total, int) or total != len(data):
        raise ValueError("DataSaudi PMI response is incomplete or has an invalid total.")

    note = (
        "Retrieved from DataSaudi's public PMI monthly data endpoint. DataSaudi attributes "
        "the series to Riyad Bank and links the S&P Global PMI release archive."
    )
    rows = []
    seen_periods: set[str] = set()
    for item in data:
        if not isinstance(item, dict):
            raise ValueError("DataSaudi PMI response contains an invalid observation.")
        period = item.get("Month")
        value = item.get(expected_measure)
        if not isinstance(period, str) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", period):
            raise ValueError(f"DataSaudi PMI response has an invalid month: {period!r}.")
        if period in seen_periods:
            raise ValueError(f"DataSaudi PMI response repeats month {period}.")
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 100:
            raise ValueError(f"DataSaudi PMI response has an invalid value for {period}.")
        seen_periods.add(period)
        year, month = (int(part) for part in period.split("-"))
        rows.append(
            make_fact(
                "KPI_01",
                year,
                month,
                value,
                "index points",
                "Purchasing Managers' Index - Monthly (DataSaudi)",
                DATASAUDI_PMI_PAGE,
                note,
            )
        )
    return sorted(rows, key=lambda row: row["PeriodStart"])


def monthly_rows(sheet, year_index: int, month_index: int):
    for row in sheet.iter_rows(values_only=True):
        if len(row) <= max(year_index, month_index):
            continue
        month = MONTHS.get(str(row[month_index] or "").strip().lower())
        if not month:
            continue
        try:
            year = int(str(row[year_index]).replace("*", "").strip())
        except (TypeError, ValueError):
            continue
        yield year, month, row


def extract_trade(content: bytes, source_url: str) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {
        key: [] for key in ("KPI_05", "KPI_07", "KPI_08", "KPI_11", "KPI_12")
    }
    for year, month, row in monthly_rows(workbook["1.1"], 0, 2):
        if row[6] is not None:
            result["KPI_05"].append(make_fact(
                "KPI_05", year, month, row[6], "SAR million",
                "GASTAT ITR 1.1 - monthly goods imports", source_url,
                "Extracted from the source workbook's monthly imports value column.",
            ))
        if row[3] is not None:
            result["KPI_12"].append(make_fact(
                "KPI_12", year, month, row[3], "SAR million",
                "GASTAT ITR 1.1 - monthly goods exports", source_url,
                "Extracted from the source workbook's monthly exports value column.",
            ))

    for year, month, row in monthly_rows(workbook["1.2"], 0, 2):
        if row[7] is not None:
            result["KPI_11"].append(make_fact(
                "KPI_11", year, month, row[7], "SAR million",
                "GASTAT ITR 1.2 - monthly oil exports", source_url,
                "Extracted from the national oil exports value column.",
            ))
        if row[3] is not None and row[11] is not None:
            total = row[3] + row[11]
            result["KPI_08"].append(make_fact(
                "KPI_08", year, month, total, "SAR million",
                "GASTAT ITR 1.2 - non-oil exports including re-exports", source_url,
                "National non-oil exports plus re-exports.",
            ))
            result.setdefault("KPI_08_Components", []).append({
                "CountryCode": "SAU",
                "KPI_ID": "KPI_08",
                "Year": year,
                "Month": month,
                "MonthName": date(year, month, 1).strftime("%B"),
                "National_NonOil_Exports_SAR_million": float(row[3]),
                "ReExports_SAR_million": float(row[11]),
                "KPI_Total_Including_ReExports_SAR_million": float(total),
                "ObservationStatus": "Reported",
                "SourceTable": "GASTAT ITR 1.2",
                "SourceURL": source_url,
            })

    sheet = workbook["7.2"]
    header_years = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    header_months = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    section_headers = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
    import_start = next(
        (
            index
            for index, value in enumerate(section_headers)
            if value and ("import" in str(value).lower() or "واردات" in str(value))
        ),
        None,
    )
    if import_start is None:
        workbook.close()
        raise ValueError("Could not locate the imports section in GASTAT ITR table 7.2")
    mapped_columns = []
    current_year = None
    for index in range(import_start, len(header_months)):
        year_value = header_years[index] if index < len(header_years) else None
        if year_value is not None:
            try:
                current_year = int(str(year_value).replace("*", "").strip())
            except ValueError:
                current_year = None
        raw_label = str(header_months[index] or "")
        label = raw_label.splitlines()[-1].strip().lower() if raw_label.strip() else ""
        month = next(
            (value for name, value in MONTHS.items() if label in {name, name[:3]}),
            None,
        )
        if current_year and month:
            mapped_columns.append((index, current_year, month))
    for row in sheet.iter_rows(min_row=7, values_only=True):
        code = str(row[0]).strip() if row and row[0] is not None else ""
        if code not in BEC_CODES:
            continue
        for index, year, month in mapped_columns:
            if index < len(row) and row[index] is not None:
                result["KPI_07"].append(
                    make_fact(
                        "KPI_07", year, month, row[index], "SAR million",
                        "GASTAT ITR 7.2 - BEC Rev.5 capital imports",
                        source_url,
                        f"Gross Fixed Capital Formation goods, BEC code {code}.",
                    )
                    | {
                        "BEC_Code": code,
                        "BEC_Description_English": (
                            str(row[56]).strip()
                            if len(row) > 56 and row[56]
                            else f"BEC Rev.5 code {code}"
                        ),
                    }
                )
    workbook.close()

    monthly_totals: dict[tuple[int, int], float] = {}
    for row in result["KPI_07"]:
        key = (row["Year"], row["Month"])
        monthly_totals[key] = monthly_totals.get(key, 0.0) + row["Value"]
    result["KPI_07_BEC"] = result["KPI_07"]
    result["KPI_07"] = [
        make_fact(
            "KPI_07", year, month, total, "SAR million",
            "GASTAT ITR 7.2 - BEC Rev.5 capital-goods aggregate",
            source_url,
            "Calculated from eight BEC Rev.5 Gross Fixed Capital Formation goods categories.",
        )
        for (year, month), total in sorted(monthly_totals.items())
    ]
    return result


def extract_august(content: bytes, source_url: str) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["1"]
    row = next(
        (
            row for row in sheet.iter_rows(min_row=5, values_only=True)
            if len(row) > 5 and str(row[2] or "").strip().lower() == "august"
        ),
        None,
    )
    workbook.close()
    if row is None:
        raise ValueError("August observation was not found in GASTAT table 1")
    year = int(str(row[0]).replace("*", "").strip())
    status = "Preliminary" if "*" in str(row[0]) else "Reported"
    return {
        "KPI_05": make_fact(
            "KPI_05", year, 8, row[5], "SAR million",
            "GASTAT non-oil trade workbook - monthly imports", source_url,
            "Latest available month from the dedicated monthly workbook.", status,
        ),
        "KPI_08": make_fact(
            "KPI_08", year, 8, row[3] + row[4], "SAR million",
            "GASTAT non-oil trade workbook - non-oil exports and re-exports", source_url,
            "National non-oil exports plus re-exports.", status,
        ),
        "KPI_08_Components": {
            "CountryCode": "SAU",
            "KPI_ID": "KPI_08",
            "Year": year,
            "Month": 8,
            "MonthName": "August",
            "National_NonOil_Exports_SAR_million": float(row[3]),
            "ReExports_SAR_million": float(row[4]),
            "KPI_Total_Including_ReExports_SAR_million": float(row[3] + row[4]),
            "ObservationStatus": status,
            "SourceTable": "GASTAT non-oil trade workbook, table 1",
            "SourceURL": source_url,
        },
    }


def extract_building_permits(content: bytes, source_url: str) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    result = []
    for row in workbook["7"].iter_rows(min_row=10, values_only=True):
        if len(row) < 13:
            continue
        try:
            year = int(row[0])
        except (TypeError, ValueError):
            continue
        for month, value in enumerate(row[1:13], start=1):
            if value is not None:
                raw_value = str(value).strip()
                observation_status = "Preliminary" if "*" in raw_value else "Reported"
                numeric_value = float(raw_value.replace(",", "").replace("*", ""))
                result.append(make_fact(
                    "KPI_13", year, month, numeric_value, "permits",
                    "GASTAT STS table 7 - issued building permits", source_url,
                    "Extracted from the official monthly time-series table.",
                    observation_status,
                ))
    workbook.close()
    return result


class MartService:
    def __init__(self, workbook_path: Path, refresh_minutes: int) -> None:
        self.workbook_path = workbook_path
        self.refresh_minutes = refresh_minutes
        self._lock = threading.RLock()
        self._refreshing = False
        self._data: dict[str, Any] = {"catalog": [], "sources": {}, "facts": [], "components": []}
        self._last_refresh: str | None = None
        self._last_error: str | None = None
        self._next_refresh: str | None = None
        self._refresh_states: dict[str, dict[str, Any]] = {}
        self.refresh()

    def refresh(self, live: bool = False) -> dict[str, Any]:
        with self._lock:
            if self._refreshing:
                return {"ok": False, "message": "A refresh is already running."}
            self._refreshing = True
        started = time.monotonic()
        try:
            data = workbook_data(self.workbook_path)
            timestamp = datetime.now(RIYADH_TZ).isoformat(timespec="seconds")
            with self._lock:
                self._data = data
                self._last_refresh = timestamp
                self._last_error = None
                counts: dict[str, int] = {}
                for fact in data["facts"]:
                    counts[fact["KPI_ID"]] = counts.get(fact["KPI_ID"], 0) + 1
                self._refresh_states = {}
                for item in data["catalog"]:
                    kpi_id = item["KPI_ID"]
                    excluded = kpi_id == "KPI_10"
                    has_data = counts.get(kpi_id, 0) > 0
                    self._refresh_states[kpi_id] = {
                        "status": "green" if has_data and not excluded else "red",
                        "label": (
                            "Excluded by request"
                            if excluded
                            else "Snapshot loaded; awaiting live extraction"
                            if has_data
                            else "No observations loaded"
                        ),
                        "note": "Live provider extraction has not run yet.",
                        "checkedAt": timestamp,
                    }
                self._next_refresh = (
                    datetime.now(RIYADH_TZ) + timedelta(minutes=self.refresh_minutes)
                ).isoformat(timespec="seconds")
            if live:
                self._extract_live_sources(timestamp)
            return {
                "ok": True,
                "lastRefresh": timestamp,
                "kpiCount": len(data["catalog"]),
                "factCount": len(self._data["facts"]),
                "durationSeconds": round(time.monotonic() - started, 3),
                "message": (
                    "DataSaudi PMI data, configured GASTAT reports, and MOF reports were fetched."
                    if live
                    else "Workbook loaded; no new provider values were fetched."
                ),
            }
        except (OSError, ValueError, KeyError) as error:
            with self._lock:
                self._last_error = str(error)
                self._next_refresh = None
            raise
        finally:
            with self._lock:
                self._refreshing = False

    def _extract_live_sources(self, timestamp: str) -> None:
        with self._lock:
            sources = dict(self._data["sources"])
            existing_facts = list(self._data["facts"])
        source_url = {
            kpi_id: sources.get(kpi_id, {}).get("ExtractedDataURL") or ""
            for kpi_id in ("KPI_05", "KPI_07", "KPI_08", "KPI_11", "KPI_12", "KPI_13")
        }
        downloaded: dict[str, bytes] = {}
        errors: dict[str, str] = {}
        extracted: dict[str, list[dict[str, Any]]] = {}
        notes: dict[str, str] = {}

        def fetch(url: str) -> bytes:
            if not url:
                raise ValueError("No direct source workbook URL is configured.")
            if url not in downloaded:
                downloaded[url] = download_bytes(url)
            return downloaded[url]

        try:
            extracted["KPI_01"] = extract_datasaudi_pmi(fetch(DATASAUDI_PMI_URL))
            notes["KPI_01"] = (
                "Extracted from DataSaudi's public PMI monthly data endpoint; the series is attributed "
                "to Riyad Bank and links the S&P Global PMI release archive."
            )
        except (OSError, ValueError, KeyError, IndexError) as error:
            errors["KPI_01"] = str(error)

        main_url = source_url["KPI_05"]
        extracted_components: list[dict[str, Any]] = []
        try:
            extracted.update(extract_trade(fetch(main_url), main_url))
            extracted_components = extracted.get("KPI_08_Components", [])
            for kpi_id in ("KPI_07", "KPI_11", "KPI_12"):
                notes[kpi_id] = "Extracted from the configured official GASTAT ITR workbook."
        except (OSError, ValueError, KeyError, IndexError) as error:
            for kpi_id in ("KPI_05", "KPI_07", "KPI_11", "KPI_12"):
                errors[kpi_id] = str(error)

        august_url = source_url["KPI_08"]
        try:
            august = extract_august(fetch(august_url), august_url)
            for kpi_id in ("KPI_05", "KPI_08"):
                rows = extracted.get(kpi_id, [])
                if rows and kpi_id not in errors:
                    rows = [row for row in rows if row["PeriodStart"] != august[kpi_id]["PeriodStart"]]
                    rows.append(august[kpi_id])
                    extracted[kpi_id] = rows
                    notes[kpi_id] = "Extracted from GASTAT ITR and the dedicated monthly trade workbook."
                    if kpi_id == "KPI_08":
                        extracted_components = [
                            row for row in extracted_components if row["Month"] != 8 or row["Year"] != 2026
                        ]
                        extracted_components.append(august["KPI_08_Components"])
                else:
                    errors.setdefault(kpi_id, "Main GASTAT ITR report was not extracted; retained snapshot data.")
        except (OSError, ValueError, KeyError, IndexError) as error:
            for kpi_id in ("KPI_05", "KPI_08"):
                if not extracted.get(kpi_id):
                    errors[kpi_id] = str(error)
                else:
                    notes[kpi_id] = f"Dedicated latest-month workbook unavailable; main ITR data loaded. {error}"

        try:
            permits_url = source_url["KPI_13"]
            extracted["KPI_13"] = extract_building_permits(fetch(permits_url), permits_url)
            notes["KPI_13"] = "Extracted from the configured official GASTAT STS workbook."
        except (OSError, ValueError, KeyError, IndexError) as error:
            errors["KPI_13"] = str(error)

        try:
            capex: list[dict[str, Any]] = []
            for report_url in mof_report_urls(download_bytes):
                capex.extend(extract_mof_capex(download_bytes(report_url), report_url))
            if not capex:
                raise ValueError("No MOF capital expenditure observations were extracted.")
            extracted["KPI_16"] = capex
            notes["KPI_16"] = "Extracted from the latest MOF quarterly budget performance report of each fiscal year."
        except (OSError, ValueError, KeyError, IndexError) as error:
            errors["KPI_16"] = str(error)

        replace_ids = {"KPI_01", "KPI_05", "KPI_07", "KPI_08", "KPI_11", "KPI_12", "KPI_13", "KPI_16"}
        successful = {
            kpi_id: rows for kpi_id, rows in extracted.items()
            if kpi_id in replace_ids and rows and kpi_id not in errors
        }
        replacement_facts = [row for rows in successful.values() for row in rows]
        with self._lock:
            self._data["facts"] = [
                row for row in existing_facts if row["KPI_ID"] not in successful
            ] + replacement_facts
            if "KPI_07" in successful:
                self._data["components"] = extracted.get("KPI_07_BEC", [])
            if "KPI_08" in successful and extracted_components:
                self._data["kpi08_components"] = extracted_components
            for item in self._data["catalog"]:
                kpi_id = item["KPI_ID"]
                if kpi_id == "KPI_10":
                    self._refresh_states[kpi_id] = {
                        "status": "red",
                        "label": "Excluded by request",
                        "note": "MISA extraction remains excluded per user instruction.",
                        "checkedAt": timestamp,
                    }
                elif kpi_id in errors:
                    self._refresh_states[kpi_id] = {
                        "status": "red",
                        "label": "Live extraction failed",
                        "note": errors[kpi_id],
                        "checkedAt": timestamp,
                    }
                elif kpi_id in successful:
                    self._refresh_states[kpi_id] = {
                        "status": "green",
                        "label": "Live extraction succeeded",
                        "note": notes.get(kpi_id, "Extracted from the configured official source workbook."),
                        "checkedAt": timestamp,
                    }
                elif kpi_id in replace_ids:
                    self._refresh_states[kpi_id] = {
                        "status": "red",
                        "label": "No live observations extracted",
                        "note": "The configured source workbook did not yield observations.",
                        "checkedAt": timestamp,
                    }
                else:
                    self._refresh_states[kpi_id] = {
                        "status": "red",
                        "label": "No live connector configured",
                        "note": "No verified automated provider connector is available for this KPI.",
                        "checkedAt": timestamp,
                    }

    def snapshot(self) -> dict[str, Any]:
        with self._lock:
            catalog = list(self._data["catalog"])
            sources = dict(self._data["sources"])
            facts = list(self._data["facts"])
            last_refresh = self._last_refresh
            last_error = self._last_error
            next_refresh = self._next_refresh
            refreshing = self._refreshing
            refresh_states = dict(self._refresh_states)

        counts: dict[str, int] = {}
        latest: dict[str, str] = {}
        for fact in facts:
            kpi_id = fact["KPI_ID"]
            counts[kpi_id] = counts.get(kpi_id, 0) + 1
            period = fact.get("PeriodEnd") or fact.get("PeriodStart") or ""
            if period and period > latest.get(kpi_id, ""):
                latest[kpi_id] = period

        kpis = []
        for item in catalog:
            kpi_id = item.get("KPI_ID")
            count = counts.get(kpi_id, 0)
            declared_status = str(item.get("ExtractionStatus") or "")
            arabic_name = item.get("KPI_Name_Arabic")
            if not arabic_name or "?" in str(arabic_name):
                arabic_name = item.get("KPI_Name_Original")
            frequency = item.get("Frequency_or_TimeFrame")
            if frequency and "?" in str(frequency) and "/" in str(frequency):
                frequency = str(frequency).rsplit("/", 1)[-1].strip()
            is_excluded = kpi_id == "KPI_10" or "Excluded" in declared_status
            state = refresh_states.get(kpi_id, {})
            status = state.get("status", "green" if count and not is_excluded else "red")
            source = sources.get(kpi_id, {})
            kpis.append(
                {
                    "id": kpi_id,
                    "name": KPI_NAMES_EN.get(kpi_id, item.get("KPI_Name_Original")),
                    "nameArabic": arabic_name,
                    "source": item.get("Source"),
                    "sourceUrl": item.get("SourceURL"),
                    "frequency": frequency,
                    "extractionStatus": declared_status,
                    "observationCount": count,
                    "latestObservation": latest.get(kpi_id),
                    "refreshStatus": status,
                    "refreshLabel": state.get("label") or (
                        "Excluded by request" if is_excluded else
                        "Workbook data loaded" if count else "No observations loaded"
                    ),
                    "refreshNote": state.get("note") or item.get("ExtractionNote") or source.get("ExtractionNote"),
                    "lastChecked": state.get("checkedAt"),
                    "observedFrom": source.get("ObservedFrom"),
                    "observedTo": source.get("ObservedTo"),
                    "sourceAccessStatus": source.get("ExtractionStatus"),
                    "liveExtractionConfigured": kpi_id in {
                        "KPI_01", "KPI_05", "KPI_07", "KPI_08", "KPI_11", "KPI_12", "KPI_13", "KPI_16"
                    },
                }
            )

        return {
            "kpis": kpis,
            "lastRefresh": last_refresh,
            "nextRefresh": next_refresh,
            "refreshIntervalMinutes": self.refresh_minutes,
            "refreshing": refreshing,
            "lastError": last_error,
            "liveExtractionConfigured": True,
            "notice": (
                "Scheduled refresh extracts the DataSaudi PMI monthly series for KPI 01, downloads official Excel "
                "reports from GASTAT for KPIs 05, 07, 08, "
                "11, 12 and 13 (fixed report URLs from the source register) and the Ministry of Finance quarterly "
                "budget performance reports for KPI 16 (report links are discovered on the MOF year pages). "
                "Other KPIs need verified provider-specific connectors."
            ),
            "factCount": len(facts),
        }

    def facts_for(self, kpi_id: str | None = None) -> list[dict[str, Any]]:
        with self._lock:
            facts = list(self._data["facts"])
        if kpi_id:
            facts = [row for row in facts if row["KPI_ID"] == kpi_id]
        return facts

    def workbook_path_for_export(self) -> Path:
        return self.workbook_path

    def export_bytes(self) -> bytes:
        facts = self.facts_for()
        with self._lock:
            components = list(self._data.get("components", []))
            export_components = list(self._data.get("kpi08_components", []))
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_names: set[str] = set()

        def add_table(sheet_name: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
            # Drop columns with no values; keep the core columns when there is no data.
            populated = [h for h in headers if any(r.get(h) is not None for r in rows)]
            columns = populated or headers[:13]
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(columns)
            for row in rows:
                sheet.append([row.get(h) for h in columns])
            if not rows:
                sheet.append([None] * len(columns))
            table_name = "tbl_" + sheet_name
            assert table_name not in used_names
            used_names.add(table_name)
            table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 2)}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
            sheet.freeze_panes = "A2"
            for index, column in enumerate(columns, start=1):
                width = max([len(str(column))] + [len(str(r.get(column))) for r in rows[:200] if r.get(column) is not None])
                sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)

        for kpi_id in KPI_NAMES_EN:
            number = kpi_id.split("_")[1]
            add_table(
                f"KPI_{number}_{KPI_SHEET_SLUGS[kpi_id]}",
                [r for r in facts if r["KPI_ID"] == kpi_id],
                list(FACT_HEADERS),
            )
        if components:
            add_table("KPI_07_BEC_Detail", components, list(components[0]))
        if export_components:
            add_table("KPI_08_Components", export_components, list(export_components[0]))
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


HTML = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Saudi KPI Data Monitor</title>
<style>
:root{font-family:Inter,Segoe UI,Arial,sans-serif;color:#162238;background:#f3f6fb}
*{box-sizing:border-box}body{margin:0}.top{background:#102844;color:#fff;padding:24px max(22px,calc((100vw - 1200px)/2))}
.top h1{margin:0;font-size:25px}.top p{margin:8px 0 0;color:#cbd6e2}.wrap{max-width:1200px;margin:24px auto;padding:0 20px}
.notice{background:#fff7df;border:1px solid #f0d58b;color:#624b08;border-radius:10px;padding:14px 16px;line-height:1.45}
.toolbar{display:flex;flex-wrap:wrap;gap:10px;margin:18px 0;align-items:center}.btn{border:0;border-radius:8px;padding:10px 14px;font-weight:650;cursor:pointer;background:#1670d2;color:#fff;text-decoration:none}
.btn.secondary{background:#e8eef5;color:#18324f}.btn:disabled{opacity:.55;cursor:wait}.stats{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px;margin:18px 0}
.stat,.panel{background:white;border:1px solid #e1e8f0;border-radius:12px;padding:16px}.stat small{color:#66758a}.stat strong{display:block;font-size:23px;margin-top:6px}
.tablebox{overflow:auto;background:#fff;border:1px solid #e1e8f0;border-radius:12px}table{width:100%;border-collapse:collapse;min-width:900px}
th,td{text-align:left;padding:12px 13px;border-bottom:1px solid #edf1f5;vertical-align:top}th{font-size:12px;text-transform:uppercase;color:#65758a;background:#f8fafc;position:sticky;top:0}
td{font-size:13px}.status{display:inline-flex;align-items:center;gap:7px;font-weight:700;white-space:nowrap}
.dot{height:10px;width:10px;border-radius:50%;display:inline-block}.green .dot{background:#168344}.red .dot{background:#d53b3b}.muted{color:#758398}
a{color:#1268ba}.foot{font-size:12px;color:#6a788b;margin:14px 2px}#message{margin-left:auto;color:#40536b}
@media(max-width:700px){.stats{grid-template-columns:1fr}.top h1{font-size:21px}}
</style>
</head>
<body>
<header class="top"><h1>Saudi KPI Data Monitor</h1>
<p>All KPIs · refresh monitoring · downloadable Excel · JSON API</p></header>
<main class="wrap">
<div id="notice" class="notice">Loading data and refresh status…</div>
<div class="toolbar">
<button id="refresh" class="btn" onclick="refreshNow()">Run live refresh</button>
<a class="btn secondary" href="/export.xlsx">Download Excel</a>
<a class="btn secondary" href="/api/kpis">KPI JSON</a>
<a class="btn secondary" href="/api/facts">Facts JSON</a>
<span id="message"></span>
</div>
<section class="stats">
<div class="stat"><small>KPIs tracked</small><strong id="kpiCount">—</strong></div>
<div class="stat"><small>KPIs with observations</small><strong id="readyCount">—</strong></div>
<div class="stat"><small>Fact observations</small><strong id="factCount">—</strong></div>
</section>
<div class="tablebox"><table><thead><tr>
<th>Refresh status</th><th>KPI</th><th>Source</th><th>Frequency</th><th>Records</th><th>Latest data</th><th>Extraction note</th>
</tr></thead><tbody id="rows"></tbody></table></div>
<p class="foot">Green means the configured source extraction succeeded, or snapshot data is loaded before the first live attempt. Red means extraction failed, is unavailable, or the KPI was excluded. A successful download does not guarantee a newer reporting period; check the latest observation date.</p>
</main>
<script>
function esc(v){return String(v??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}
async function load(){
 const r=await fetch('/api/status'); const d=await r.json();
 document.getElementById('notice').textContent=d.notice+(d.lastRefresh?' Last refresh: '+d.lastRefresh+'.':'');
 document.getElementById('kpiCount').textContent=d.kpis.length;
 document.getElementById('readyCount').textContent=d.kpis.filter(k=>k.observationCount>0).length;
 document.getElementById('factCount').textContent=d.factCount;
 document.getElementById('rows').innerHTML=d.kpis.map(k=>`<tr>
 <td><span class="status ${esc(k.refreshStatus)}"><i class="dot"></i>${esc(k.refreshLabel)}</span></td>
 <td><strong>${esc(k.id)}</strong><br>${esc(k.name)}<br><span class="muted">${esc(k.nameArabic)}</span></td>
 <td>${k.sourceUrl?`<a href="${esc(k.sourceUrl)}" target="_blank" rel="noreferrer">${esc(k.source||'Source')}</a>`:esc(k.source||'Not specified')}</td>
 <td>${esc(k.frequency)}</td><td>${Number(k.observationCount).toLocaleString()}</td><td>${esc(k.latestObservation||'—')}</td>
 <td>${esc(k.refreshNote||k.extractionStatus||'No extraction details')}</td></tr>`).join('');
}
async function refreshNow(){
 const b=document.getElementById('refresh'),m=document.getElementById('message');b.disabled=true;m.textContent='Fetching configured source reports…';
 try{const r=await fetch('/api/refresh',{method:'POST'});const d=await r.json();if(!r.ok)throw new Error(d.error||'Refresh failed');m.textContent=d.message;await load();}
 catch(e){m.textContent=e.message}finally{b.disabled=false;setTimeout(()=>m.textContent='',6000)}
}
load();setInterval(load,30000);
</script>
</body></html>"""


class DashboardHandler(BaseHTTPRequestHandler):
    service: MartService

    def log_message(self, format: str, *args: Any) -> None:
        print(f"{self.address_string()} - {format % args}")

    def send_json(self, payload: Any, status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            body = HTML.encode("utf-8")
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        elif parsed.path in {"/api/status", "/api/kpis"}:
            self.send_json(self.service.snapshot())
        elif parsed.path == "/api/facts":
            query = parse_qs(parsed.query)
            kpi_id = query.get("kpi_id", [None])[0]
            self.send_json({"count": len(self.service.facts_for(kpi_id)), "data": self.service.facts_for(kpi_id)})
        elif parsed.path.startswith("/api/kpis/"):
            kpi_id = parsed.path.rsplit("/", 1)[-1]
            result = next((k for k in self.service.snapshot()["kpis"] if k["id"] == kpi_id), None)
            if result is None:
                self.send_json({"error": "Unknown KPI id"}, HTTPStatus.NOT_FOUND)
            else:
                facts = self.service.facts_for(kpi_id)
                self.send_json({"kpi": result, "count": len(facts), "data": facts})
        elif parsed.path == "/export.xlsx":
            self.send_excel()
        else:
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:
        if urlparse(self.path).path != "/api/refresh":
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
            return
        try:
            self.send_json(self.service.refresh(live=True))
        except (OSError, ValueError, KeyError) as error:
            self.send_json({"error": str(error)}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def send_excel(self) -> None:
        try:
            content = self.service.export_bytes()
        except (OSError, ValueError, KeyError) as error:
            self.send_json({"error": f"Excel export unavailable: {error}"}, HTTPStatus.INTERNAL_SERVER_ERROR)
            return
        self.send_response(HTTPStatus.OK)
        self.send_header(
            "Content-Type",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.send_header("Content-Disposition", 'attachment; filename="Saudi_KPI_Extract.xlsx"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)


def scheduler(service: MartService, stop: threading.Event) -> None:
    while not stop.wait(service.refresh_minutes * 60):
        try:
            result = service.refresh(live=True)
            print(f"Scheduled workbook refresh: {result}")
        except (OSError, ValueError, KeyError) as error:
            print(f"Scheduled workbook refresh failed: {error}")


def main() -> None:
    service = MartService(WORKBOOK_PATH, REFRESH_MINUTES)
    try:
        print(f"Initial live extraction: {service.refresh(live=True)}")
    except (OSError, ValueError, KeyError) as error:
        print(f"Initial live extraction failed: {error}")
    DashboardHandler.service = service
    stop = threading.Event()
    threading.Thread(target=scheduler, args=(service, stop), daemon=True).start()
    server = ThreadingHTTPServer((HOST, PORT), DashboardHandler)
    print(f"Dashboard: http://{HOST}:{PORT}")
    print(f"Workbook: {WORKBOOK_PATH}")
    print(f"Scheduled live refresh: every {REFRESH_MINUTES} minutes")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("Shutting down dashboard.")
    finally:
        stop.set()
        server.server_close()


if __name__ == "__main__":
    main()
