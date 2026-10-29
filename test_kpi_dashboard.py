import json
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook

from kpi_dashboard import (
    DashboardHandler,
    MartService,
    extract_datasaudi_pmi,
    extract_mof_capex,
    extract_sama_pos_text,
    sama_pos_report_urls,
)


def mof_workbook(title: str, capex_values: list[float]) -> bytes:
    workbook = Workbook()
    stale = workbook.active
    stale.title = "Expenditures"
    stale.append(["x", "Actual expenditures in Q1 of FY 2023"])
    stale.append(["الأصول", 1.0, "Non-financial Assets (CAPEX)"])
    sheet = workbook.create_sheet("Expenditures ")
    sheet.append(["المصروفات", "Expenditures"])
    sheet.append([None, title])
    sheet.append(["الأصول", *capex_values, "Non-financial Assets (CAPEX)"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class MofCapexTests(unittest.TestCase):
    def test_h1_report_yields_two_quarters_from_current_sheet(self) -> None:
        rows = extract_mof_capex(
            mof_workbook("Actual Expenditures in H1 of FY 2026", [89659.0, 46226.0, 43433.0]), "https://www.mof.gov.sa/x"
        )
        self.assertEqual([(r["PeriodStart"], r["Value"]) for r in rows],
                         [("2026-01-01", 43433.0), ("2026-04-01", 46226.0)])
        self.assertEqual(rows[1]["PeriodEnd"], "2026-06-30")

    def test_unreconciled_totals_are_rejected(self) -> None:
        with self.assertRaises(ValueError):
            extract_mof_capex(mof_workbook("Actual expenditures until Q3 of FY 2024", [100.0, 1.0, 2.0, 3.0]), "u")


class DataSaudiPmiTests(unittest.TestCase):
    def test_extracts_monthly_pmi_facts(self) -> None:
        payload = {
            "page": {"total": 2},
            "columns": ["Month ID", "Month", "Purchasing Manager Index"],
            "data": [
                {"Month ID": 202609, "Month": "2026-09", "Purchasing Manager Index": 55.26107857},
                {"Month ID": 202608, "Month": "2026-08", "Purchasing Manager Index": 53.75571254},
            ],
        }
        rows = extract_datasaudi_pmi(json.dumps(payload).encode("utf-8"))
        self.assertEqual([row["PeriodStart"] for row in rows], ["2026-08-01", "2026-09-01"])
        self.assertEqual(rows[-1]["Value"], 55.26107857)
        self.assertEqual(rows[-1]["Unit"], "index points")
        self.assertEqual(rows[-1]["SourceURL"], "https://datasaudi.sa/en/#purchasing-manager-index")

    def test_rejects_incomplete_api_response(self) -> None:
        payload = {
            "page": {"total": 2},
            "columns": ["Month", "Purchasing Manager Index"],
            "data": [{"Month": "2026-09", "Purchasing Manager Index": 55.3}],
        }
        with self.assertRaisesRegex(ValueError, "incomplete"):
            extract_datasaudi_pmi(json.dumps(payload).encode("utf-8"))


class SamaPosTests(unittest.TestCase):
    def test_finds_report_pdf_links_in_sharepoint_list_data(self) -> None:
        html = r'''{"FileRef":"\u002fen-US\u002fStatistics\u002fIndices\u002fPOS_EN\u002fWeekly_Points_of_Sale_Transactions_Report.pdf"}'''
        urls = sama_pos_report_urls(lambda _: html.encode("utf-8"))
        self.assertEqual(len(urls), 1)
        self.assertTrue(urls[0].endswith("/POS_EN/Weekly_Points_of_Sale_Transactions_Report.pdf"))

    def test_extracts_weekly_count_and_value_as_separate_rows(self) -> None:
        text = """
        27 Sep,26 - 03 Oct,26 04 Oct,26 - 10 Oct,26 11 Oct,26 - 17 Oct,26 18 Oct,26 - 24 Oct,26
        Number of Transactions Value of Transactions
        Total 273,924 16,507,357 253,359 13,525,658 246,421 12,602,805 229,116 11,779,845 -7.0 -6.5
        """
        rows = extract_sama_pos_text(text, "https://www.sama.gov.sa/report.pdf")
        self.assertEqual(len(rows), 8)
        latest = [row for row in rows if row["PeriodStart"] == "2026-10-18"]
        self.assertEqual({row["SourceTable"] for row in latest}, {
            "SAMA POS total transaction count",
            "SAMA POS total transaction value",
        })
        count = next(row for row in latest if row["Unit"] == "thousand transactions")
        value = next(row for row in latest if row["Unit"] == "SAR thousand")
        self.assertEqual((count["Value"], count["PeriodEnd"]), (229116.0, "2026-10-24"))
        self.assertEqual(value["Value"], 11779845.0)

    def test_rejects_report_without_total_or_periods(self) -> None:
        with self.assertRaisesRegex(ValueError, "four weekly periods or total"):
            extract_sama_pos_text("Weekly Points of Sale Transactions", "u")


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    catalog = workbook.active
    catalog.title = "KPI_Catalog"
    catalog.append(
        [
            "KPI_ID",
            "KPI_Name_Arabic",
            "KPI_Name_Original",
            "Source",
            "SourceURL",
            "Frequency_or_TimeFrame",
            "ExtractionStatus",
            "ExtractionNote",
        ]
    )
    catalog.append(["KPI_05", "الواردات", "Imports", "GASTAT", "", "Monthly", "Extracted", "Fixture"])
    catalog.append(["KPI_10", "التراخيص", "Licences", "MISA", "", "Monthly", "Excluded", "Excluded by request"])

    sources = workbook.create_sheet("Source_Register")
    sources.append(["KPI_ID", "ObservedFrom", "ObservedTo", "ExtractionStatus", "ExtractionNote"])
    sources.append(["KPI_05", "2025-01", "2025-01", "Extracted", "Fixture"])
    sources.append(["KPI_10", None, None, "Excluded", "Excluded by request"])

    facts = workbook.create_sheet("Fact_All")
    facts.append(
        [
            "CountryCode",
            "KPI_ID",
            "PeriodStart",
            "PeriodEnd",
            "Year",
            "Month",
            "MonthName",
            "Value",
            "Unit",
            "SourceTable",
            "ObservationStatus",
            "SourceURL",
            "SourceNote",
            "Open",
            "High",
            "Low",
            "Close",
            "VolumeTraded",
            "TurnoverSAR",
            "TradeCount",
        ]
    )
    facts.append(
        [
            "SAU",
            "KPI_05",
            "2025-01-01",
            "2025-01-31",
            2025,
            1,
            "January",
            12.5,
            "SAR million",
            "Test table",
            "Reported",
            "",
            "",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    workbook.save(path)


class MartServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "test.xlsx"
        make_workbook(self.workbook_path)
        self.service = MartService(self.workbook_path, refresh_minutes=60)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_status_includes_all_kpis_and_distinguishes_missing_data(self) -> None:
        result = self.service.snapshot()
        self.assertEqual(len(result["kpis"]), 2)
        self.assertEqual(result["factCount"], 1)
        self.assertEqual(result["kpis"][0]["refreshStatus"], "green")
        self.assertEqual(result["kpis"][0]["name"], "Total goods imports")
        self.assertEqual(result["kpis"][0]["nameArabic"], "الواردات")
        self.assertEqual(result["kpis"][0]["frequency"], "Monthly")
        self.assertEqual(result["kpis"][1]["refreshLabel"], "Excluded by request")
        self.assertEqual(result["kpis"][1]["refreshStatus"], "red")

    def test_facts_can_be_filtered(self) -> None:
        self.assertEqual(len(self.service.facts_for("KPI_05")), 1)
        self.assertEqual(self.service.facts_for("KPI_10"), [])

    def test_refresh_reloads_updated_workbook(self) -> None:
        result = self.service.refresh()
        self.assertTrue(result["ok"])
        self.assertEqual(result["factCount"], 1)
        self.assertIsNotNone(self.service.snapshot()["lastRefresh"])

    def test_excel_export_has_one_named_table_per_kpi(self) -> None:
        workbook = load_workbook(BytesIO(self.service.export_bytes()), data_only=True)
        kpi_sheets = [name for name in workbook.sheetnames if name.startswith("KPI_")]
        self.assertEqual(len(kpi_sheets), 16)
        self.assertNotIn("Refresh_Status", workbook.sheetnames)
        self.assertNotIn("Fact_All", workbook.sheetnames)
        sheet = workbook["KPI_05_Total_Goods_Imports"]
        self.assertEqual(list(sheet.tables), ["tbl_KPI_05_Total_Goods_Imports"])
        self.assertEqual(sheet.max_row, 2)


class DashboardApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        path = Path(self.temp_dir.name) / "test.xlsx"
        make_workbook(path)
        DashboardHandler.service = MartService(path, refresh_minutes=60)
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), DashboardHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        self.temp_dir.cleanup()

    def test_kpi_api_returns_json(self) -> None:
        with urlopen(f"{self.base_url}/api/kpis") as response:
            self.assertIn("charset=utf-8", response.headers["Content-Type"])
            raw = response.read()
            self.assertIn("الواردات".encode("utf-8"), raw)
            body = json.loads(raw.decode("utf-8"))
        self.assertEqual(len(body["kpis"]), 2)

    def test_filtered_fact_api_returns_rows(self) -> None:
        with urlopen(f"{self.base_url}/api/facts?kpi_id=KPI_05") as response:
            body = json.loads(response.read())
        self.assertEqual(body["count"], 1)
        self.assertEqual(body["data"][0]["Value"], 12.5)

    def test_unknown_kpi_returns_404(self) -> None:
        with self.assertRaises(HTTPError) as error:
            urlopen(f"{self.base_url}/api/kpis/KPI_99")
        self.assertEqual(error.exception.code, 404)

    def test_refresh_endpoint_returns_status(self) -> None:
        request = Request(f"{self.base_url}/api/refresh", method="POST")
        with urlopen(request) as response:
            body = json.loads(response.read())
        self.assertTrue(body["ok"])
        self.assertEqual(body["kpiCount"], 2)

    def test_excel_download_returns_workbook(self) -> None:
        with urlopen(f"{self.base_url}/export.xlsx") as response:
            self.assertEqual(response.status, 200)
            self.assertIn("spreadsheetml", response.headers["Content-Type"])
            workbook = load_workbook(BytesIO(response.read()), read_only=True, data_only=True)
        self.assertIn("KPI_05_Total_Goods_Imports", workbook.sheetnames)
        self.assertEqual(workbook["KPI_05_Total_Goods_Imports"].max_row, 2)


if __name__ == "__main__":
    unittest.main()
